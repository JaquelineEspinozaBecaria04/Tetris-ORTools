# excel_layout.py
import pandas as pd
import numpy as np
import io
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from typing import List, Dict, Any

# Constantes del Notebook
NUMERO_CHIPS = 2
NUMERO_CORES = 20
COL_INFO = ['# S. TOTAL', '# S. ZONA', 'SERVIDOR']
INFRA_COLOR = 'FBE2D5' # Color sin '#'

def generar_diccionario(df_solucion: pd.DataFrame) -> Dict[str, List[Any]]:
    """Crea el diccionario de lookup {VM: [Length, Color]}."""
    df_diccionario = df_solucion[['VM', 'Length', 'Color']]
    df_unique = df_diccionario.drop_duplicates(subset="VM")
    diccionario = {fila.VM: [fila.Length, fila.Color] for fila in df_unique.itertuples(index=False)}
    
    diccionario['INFRA'] = [3, f"#{INFRA_COLOR}"] # Guardar con #
    diccionario.setdefault(' ', [1, '#FFFFFF'])
    
    # Asegurarse de que los valores nulos o NaN no causen problemas
    diccionario = {k: v for k, v in diccionario.items() if pd.notna(k)}

    for az in df_solucion['AZ'].unique():
        if pd.notna(az) and az not in diccionario:
            diccionario[az] = [1, '#696969']

    return diccionario

def generar_matriz(df_solucion: pd.DataFrame) -> List[List[Any]]:
    """Pivota el DataFrame largo a una matriz ancha basada en la posición (Start)."""
    # Validación: Salir si el DF está vacío o la columna HOST no existe/está vacía
    if df_solucion.empty or 'HOST' not in df_solucion.columns or df_solucion['HOST'].isnull().all():
        return []
        
    total_hosts = int(df_solucion['HOST'].max())
    matriz = [[np.nan for _ in range((NUMERO_CORES * NUMERO_CHIPS) + len(COL_INFO) + 1)] for _ in range(total_hosts)]
    
    df_piezas = df_solucion.dropna(subset=['Start'])

    for fila in df_piezas.itertuples(index=False):
        try:
            start_pos = int(fila.Start)
            host_idx = int(fila.Nuevo_HOST) # 0-based index
        except (ValueError, TypeError):
            continue 

        if host_idx >= len(matriz):
            print(f"Advertencia: Índice de host {host_idx} fuera de rango. Omitiendo fila.")
            continue
            
        matriz[host_idx][start_pos] = fila.VM
        
        matriz[host_idx][-4] = fila.AZ
        matriz[host_idx][-3] = fila.HOST # # S. TOTAL
        matriz[host_idx][-2] = fila.HOST_RELATIVO # # S. ZONA
        
        # 'Host_str' (SERVIDOR) puede no existir si df_layout falló, poner fallback
        matriz[host_idx][-1] = getattr(fila, 'Host_str', f"Host {fila.HOST}") # SERVIDOR (string)
    
    return matriz

def definir_columnas(col_info: List[str]) -> List[str]:
    """Define los nombres de las columnas para la matriz ancha."""
    chip_cols = [f'chip_{i}_{j}' for i in range(1, NUMERO_CHIPS + 1) for j in range(1, NUMERO_CORES + 1)]
    return chip_cols + ['AZ'] + col_info

def obtener_df_final(matriz: List[List[Any]], columnas: List[str]) -> pd.DataFrame:
    """Convierte la matriz en el DataFrame final, insertando filas de separación por AZ."""
    if not matriz:
        return pd.DataFrame(columns=columnas)
        
    df = pd.DataFrame(matriz, columns=columnas)
    
    df = df.dropna(how='all').reset_index(drop=True)
    if df.empty:
        return df

    cambios = df["AZ"] != df["AZ"].shift()
    filas_cambio = df.index[cambios].tolist()
    
    for idx in reversed(filas_cambio[1:]):
        df = pd.concat([
            df.iloc[:idx],
            pd.DataFrame([{}], columns=df.columns), # Fila vacía
            df.iloc[idx:]
        ], ignore_index=True)
    
    info_cols = COL_INFO
    chip_1_cols = [c for c in columnas if c.startswith('chip_1_')]
    chip_2_cols = [c for c in columnas if c.startswith('chip_2_')]
    az_col = ['AZ']
    
    final_order = info_cols + chip_1_cols + az_col + chip_2_cols
    final_order_existentes = [c for c in final_order if c in df.columns]
    # Asegurarse de que todas las columnas existan antes de reordenar
    missing_cols = set(final_order_existentes) - set(df.columns)
    if missing_cols:
        print(f"Advertencia: Faltan columnas en df_final: {missing_cols}")
        final_order_existentes = [c for c in final_order_existentes if c in df.columns]

    df = df[final_order_existentes]
    
    return df

def aplicar_estilos_excel(writer: pd.ExcelWriter, df: pd.DataFrame, diccionario: Dict[str, List[Any]], sitio: str):
    """
    Aplica el formato del notebook (fusionar celdas, colores, etc.) 
    usando openpyxl.
    """
    if df.empty:
        return # No hay nada que estilizar

    workbook = writer.book
    ws = writer.sheets['Layout']

    # Definir Estilos Base
    fuente = Font(name="Aptos Narrow", size=11)
    fuente_servidor = Font(name="Aptos Narrow", size=11, bold=True)
    fuente_titulo = Font(name="Aptos Narrow", size=16, bold=True, color="FF0000")
    fuente_az = Font(name="Aptos Narrow", size=14, bold=True)
    
    alineacion = Alignment(horizontal="center", vertical="center")
    alineacion_az = Alignment(horizontal="center", vertical="center", textRotation=90)

    borde = Side(style="thin", color="000000")
    bordes_completos = Border(left=borde, right=borde, top=borde, bottom=borde)
    
    # Rellenos (colores sin #)
    relleno_titulo = PatternFill(start_color='CAEDFB', end_color='CAEDFB', fill_type="solid")
    relleno_header = PatternFill(start_color=INFRA_COLOR, end_color=INFRA_COLOR, fill_type="solid")
    relleno_info_cols = PatternFill(start_color=INFRA_COLOR, end_color=INFRA_COLOR, fill_type="solid")
    relleno_az_vertical = PatternFill(start_color='B5E6A2', end_color='B5E6A2', fill_type="solid")
    relleno_infra_fila = PatternFill(start_color='000000', end_color='000000', fill_type="solid")

    filas_datos_offset = 2 # Título + Header
    max_row = df.shape[0] + filas_datos_offset
    max_col = df.shape[1]
    
    # Aplicar Estilos Generales
    for fila in range(1, max_row + 1): 
        for col in range(1, max_col + 1):
            celda = ws.cell(row=fila, column=col)
            celda.font = fuente
            celda.alignment = alineacion
            celda.border = bordes_completos

    # Estilo de Título
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    celda_titulo = ws.cell(row=1, column=1)
    celda_titulo.value = f"TETRIS ACTUAL {sitio}"
    celda_titulo.font = fuente_titulo
    celda_titulo.alignment = alineacion
    celda_titulo.fill = relleno_titulo

    # Estilo de Header
    col_chip_1_start = -1
    col_chip_2_end = -1
    try:
        col_chip_1_start = df.columns.get_loc('chip_1_1') + 1
        col_chip_2_end = df.columns.get_loc(f'chip_2_{NUMERO_CORES}') + 1
    except KeyError:
        print("Advertencia: No se encontraron columnas 'chip_1_1' o 'chip_2_20'.")
        # Buscar la última columna de chip como fallback
        cols_chip_list = [c for c in df.columns if c.startswith("chip_")]
        if cols_chip_list:
            col_chip_2_end = df.columns.get_loc(cols_chip_list[-1]) + 1


    for col in range(1, max_col + 1):
        celda = ws.cell(row=2, column=col)
        celda.fill = relleno_header
        valor_celda = celda.value
        if isinstance(valor_celda, str) and valor_celda.startswith("chip_"):
            try:
                celda.value = int(valor_celda.split("_")[-1])
            except:
                pass 

    # Estilo Columnas de Info
    col_info_end = len(COL_INFO)
    servidor_col_idx = -1
    try:
        servidor_col_idx = df.columns.get_loc('SERVIDOR') + 1
    except KeyError:
        print("Advertencia: No se encontró la columna SERVIDOR")

    for fila in range(filas_datos_offset + 1, max_row + 1):
        for col in range(1, col_info_end + 1):
            celda = ws.cell(row=fila, column=col)
            celda.fill = relleno_info_cols
            if col == servidor_col_idx:
                celda.font = fuente_servidor

    # Fusionar y Colorear VMs
    cols_chip = [c for c in df.columns if c.startswith("chip_")]
    formatos_color = {}

    for idx, fila_df in df.iterrows():
        fila_excel = idx + filas_datos_offset + 1 
        
        for col_name in cols_chip:
            vm = fila_df[col_name]
            if pd.isna(vm):
                continue
            
            vm_info = diccionario.get(vm, [1, '#FFFFFF']) 
            vm_len, vm_color = vm_info[0], vm_info[1]

            vm_len = int(vm_len) if pd.notna(vm_len) else 1
            if vm_len <= 0:
                vm_len = 1

            if not isinstance(vm_color, str) or not vm_color.startswith('#') or len(vm_color) != 7:
                vm_color = '#AB63FA' # Color por defecto

            color_hex = vm_color[1:] # Quitar #
            
            if color_hex not in formatos_color:
                formatos_color[color_hex] = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
            
            relleno_vm = formatos_color[color_hex]

            col_excel_start = df.columns.get_loc(col_name) + 1
            col_excel_end = col_excel_start + vm_len - 1
            
            # Asegurar que no se fusione más allá de la última columna de chip
            if col_chip_2_end != -1:
                 col_excel_end = min(col_excel_end, col_chip_2_end) 

            if col_excel_start <= col_excel_end:
                try:
                    ws.merge_cells(start_row=fila_excel, start_column=col_excel_start, 
                                   end_row=fila_excel, end_column=col_excel_end)
                except Exception as e:
                    print(f"Error al fusionar celdas: {e} en fila {fila_excel}")
                    pass 
            
            ws.cell(row=fila_excel, column=col_excel_start).fill = relleno_vm

    # Colorear Filas 'INFRA' (filas vacías)
    indices_nulos = df.index[df.isnull().all(axis=1)].tolist()
    
    for fila_idx in indices_nulos:
        fila_excel = fila_idx + filas_datos_offset + 1
        ws.merge_cells(start_row=fila_excel, start_column=1, end_row=fila_excel, end_column=max_col)
        celda_infra = ws.cell(row=fila_excel, column=1)
        celda_infra.value = "INFRA"
        celda_infra.fill = relleno_infra_fila
        celda_infra.font = Font(color="FFFFFF") # Letra blanca

    # Columna Vertical 'AZ'
    col_az_idx = -1
    try:
        col_az_idx = df.columns.get_loc('AZ') + 1
    except KeyError:
        print("Advertencia: No se encontró la columna AZ")

    if col_az_idx != -1:
        fila_inicio_bloque = filas_datos_offset + 1 # Fila 3
        
        bloques_az = indices_nulos + [df.shape[0]]
        
        for fila_idx in bloques_az:
            fila_fin_bloque = fila_idx + filas_datos_offset 
            
            if fila_inicio_bloque <= fila_fin_bloque:
                try:
                    ws.merge_cells(start_row=fila_inicio_bloque, start_column=col_az_idx,
                                   end_row=fila_fin_bloque, end_column=col_az_idx)
                except Exception as e:
                     print(f"Error al fusionar AZ: {e}")
                     pass 

                celda_az = ws.cell(row=fila_inicio_bloque, column=col_az_idx)
                celda_az.font = fuente_az
                celda_az.alignment = alineacion_az
                celda_az.fill = relleno_az_vertical
                
            fila_inicio_bloque = fila_fin_bloque + 2 # +1 por la fila nula, +1 para la siguiente

    # Ajustar anchos
    try:
        ws.column_dimensions[get_column_letter(df.columns.get_loc('SERVIDOR') + 1)].width = 30
        ws.column_dimensions[get_column_letter(df.columns.get_loc('# S. TOTAL') + 1)].width = 10
        ws.column_dimensions[get_column_letter(df.columns.get_loc('# S. ZONA') + 1)].width = 10
    except KeyError:
        pass # Ignorar si las columnas no existen


def generate_excel_bytes(df_from_layout: pd.DataFrame, sitio: str) -> io.BytesIO:
    """
    Función principal que toma el DataFrame del solver y genera
    el archivo Excel estilizado en memoria.
    """
    
    df_solucion = df_from_layout.copy()
    
    # Renombrar columnas para que coincidan con la lógica del notebook
    df_solucion.rename(columns={
        'Host': 'HOST_RELATIVO', # 'Host' (local) pasa a ser 'HOST_RELATIVO'
        'host_chip': 'Host_str'  # 'host_chip' (etiqueta Y) pasa a ser 'Host_str' (usado como SERVIDOR)
    }, inplace=True)
    
    # Crear 'Nuevo_HOST' (índice 0-based)
    if 'HOST' in df_solucion.columns:
        df_solucion['Nuevo_HOST'] = df_solucion['HOST'] - 1
    else:
        df_solucion['Nuevo_HOST'] = 0 # Fallback

    df_solucion = df_solucion.sort_values(by=["HOST", "HOST_RELATIVO", "Chip", "Start"])
    
    diccionario = generar_diccionario(df_solucion)
    matriz = generar_matriz(df_solucion)
    columnas = definir_columnas(COL_INFO)
    df_final = obtener_df_final(matriz, columnas)

    output = io.BytesIO()
    # Usar 'openpyxl' como engine
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_final.to_excel(writer, sheet_name='Layout', index=False, startrow=1, header=True)
        aplicar_estilos_excel(writer, df_final, diccionario, sitio)
    
    output.seek(0)
    return output