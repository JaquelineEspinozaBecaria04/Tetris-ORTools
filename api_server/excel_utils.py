# backend/excel_utils.py
import pandas as pd
import numpy as np
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from typing import List, Dict, Any

# Constantes
NUMERO_CHIPS = 2
NUMERO_CORES = 20 # 17 usables + 3 infra
col_info = ['# S. TOTAL', '# S. ZONA', 'SERVIDOR']

def generar_excel_visual_to(df_original: pd.DataFrame, site_name: str = "Optimizado") -> bytes:
    """Genera el archivo Excel visual con el estilo exacto del notebook."""
    df = df_original.copy()
    diccionario_colores = generar_diccionario_colores(df)
    matriz = generar_matriz_con_length(df)
    
    # Definir columnas iniciales
    cols = [f'chip_{i}_{j}' for i in range(1, NUMERO_CHIPS + 1) for j in range(1, NUMERO_CORES + 1)] + ['AZ'] + col_info
    
    # Obtener DF final ordenado y con separadores
    df_excel = obtener_df_final(matriz, cols)
    
    # Generar bytes con estilo avanzado
    return generar_excel_bytes_avanzado(site_name, df_excel, diccionario_colores)

def generar_diccionario_colores(df: pd.DataFrame) -> Dict[str, str]:
    temp = df[['VM', 'Color']].drop_duplicates(subset='VM')
    dicc = {row.VM: row.Color for row in temp.itertuples(index=False)}
    dicc['INFRA'] = '#000000'
    for vm in dicc:
        if pd.isna(dicc[vm]): dicc[vm] = "#808080"
    return dicc

def generar_matriz_con_length(df: pd.DataFrame) -> List[List[Any]]:
    # Ordenamos por HOST global
    if 'HOST' in df.columns:
        df_sorted = df.sort_values(by='HOST')
        unique_hosts = df_sorted['HOST'].unique()
    else:
        unique_hosts = range(len(df))

    host_to_idx = {h: i for i, h in enumerate(unique_hosts)}
    
    total_cols = (NUMERO_CORES * NUMERO_CHIPS) + len(col_info) + 1
    matriz = [[np.nan for _ in range(total_cols)] for _ in range(len(unique_hosts))]

    for row in df.itertuples(index=False):
        host_val = getattr(row, 'HOST', None)
        idx_fila = host_to_idx.get(host_val)
        if idx_fila is None: continue

        try:
            idx_col = int(row.Start)
        except:
            continue
            
        # Guardamos Tupla (Nombre, Longitud)
        matriz[idx_fila][idx_col] = (row.VM, row.Length)

        # Metadatos
        matriz[idx_fila][-4] = row.AZ
        matriz[idx_fila][-3] = row.HOST       # # S. TOTAL
        host_relativo = getattr(row, 'HOST_RELATIVO', getattr(row, 'Host', 0))
        matriz[idx_fila][-2] = host_relativo  # # S. ZONA
        
        host_chip_str = getattr(row, 'host_chip', f"{row.AZ} - {host_relativo}")
        matriz[idx_fila][-1] = host_chip_str

    return matriz

def obtener_df_final(matriz: List[List[Any]], columnas: List[str]) -> pd.DataFrame:
    df = pd.DataFrame(matriz, columns=columnas)
    if 'SERVIDOR' in df.columns:
        df['SERVIDOR'] = df['SERVIDOR'].replace(' ', 'INFRA')
    
    if not df.empty and 'AZ' in df.columns:
        df_az = df['AZ'].astype(str)
        cambios = df_az != df_az.shift()
        indices = df.index[cambios].tolist()[1:]
        for idx in reversed(indices):
            vacio = pd.DataFrame([{}], columns=df.columns)
            df = pd.concat([df.iloc[:idx], vacio, df.iloc[idx:]], ignore_index=True)

    cols = df.columns
    try:
        # Reordenar: [INFO] + [CHIP 1] + [AZ] + [CHIP 2]
        cols_info_idx = cols[-len(col_info):].tolist()
        cols_az_idx = [cols[-(len(col_info)+1)]]
        cols_chips_idx = cols[:- (len(col_info)+1)].tolist()
        
        new_order = cols_info_idx + cols_chips_idx[:NUMERO_CORES] + cols_az_idx + cols_chips_idx[NUMERO_CORES:]
        df = df[new_order]
    except:
        pass
    return df

def generar_excel_bytes_avanzado(sitio: str, df: pd.DataFrame, diccionario_colores: Dict[str, str]) -> bytes:
    """
    Implementación portada del Notebook Gráfico_actual.ipynb
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Tetris"

    # Ajustes de filas/cols para openpyxl (1-based)
    filas_recorridas = 3 # Título (1) + Header (1) + Offset (1) -> Datos inician en 3? No, en notebook inician en 3
    
    # 1. Título
    ws.cell(row=1, column=1, value=f"TETRIS OPTIMIZADO - {sitio}")

    # 2. Headers (Fila 2)
    for c_idx, col_name in enumerate(df.columns.tolist(), 1):
        val = col_name
        if "chip_" in col_name:
            val = int(col_name.split("_")[-1])
        ws.cell(row=2, column=c_idx, value=val)

    # 3. Escribir Datos (Fila 3 en adelante)
    # Iteramos sobre el DF. Si es tupla (VM, Len), escribimos VM.
    for r_idx, row_data in enumerate(df.itertuples(index=False), 3):
        for c_idx, value in enumerate(row_data, 1):
            if isinstance(value, tuple):
                ws.cell(row=r_idx, column=c_idx, value=value[0])
            else:
                ws.cell(row=r_idx, column=c_idx, value=value)

    # 4. ESTILOS GLOBALES (Aptos Narrow 11, Centrado, Bordes)
    fuente = Font(name="Aptos Narrow", size=11)
    alineacion = Alignment(horizontal="center", vertical="center")
    borde = Side(style="thin", color="000000")
    borde_obj = Border(left=borde, right=borde, top=borde, bottom=borde)

    max_row = df.shape[0] + 2
    max_col = df.shape[1]

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = fuente
            cell.alignment = alineacion
            cell.border = borde_obj

    # 5. FUSIÓN DE CELDAS Y COLORES (Lógica del Notebook adaptada a Tuples)
    cols_list = df.columns.tolist()
    
    for i, row in enumerate(df.itertuples(index=False)):
        r_excel = i + 3 # Datos inician en fila 3
        
        # Detectar fila separadora (INFRA)
        # En el notebook: if all null -> Merge total -> Black -> INFRA
        # Aquí detectamos si la fila está vacía en el DF
        es_vacia = True
        for val in row:
            if isinstance(val, tuple) or (isinstance(val, (str, int, float)) and pd.notna(val)):
                es_vacia = False
                break
        
        if es_vacia:
    # --- INICIO CORRECCIÓN FILA NEGRA ---
            ws.merge_cells(start_row=r_excel, end_row=r_excel, start_column=1, end_column=max_col)
            
            # Value vacío ("") para que no salga texto
            cell = ws.cell(row=r_excel, column=1, value="") 
            
            # Relleno Negro
            cell.fill = PatternFill("solid", fgColor="000000")
            
            # Bordes negros para eliminar líneas blancas internas
            negro_side = Side(style="thin", color="000000")
            cell.border = Border(left=negro_side, right=negro_side, top=negro_side, bottom=negro_side)            
            continue

        # Procesar Chips (Fusión y Color)
        for j, val in enumerate(row):
            col_name = cols_list[j]
            if "chip_" not in col_name: continue
            
            if isinstance(val, tuple):
                vm_name, length = val
                length = int(length)
                
                # Obtener color
                hex_color = diccionario_colores.get(str(vm_name), "808080").replace("#", "")
                
                # Calcular rango de fusión
                c_start = j + 1
                c_end = c_start + length - 1
                
                if c_end <= max_col:
                    if length > 1:
                        try:
                            ws.merge_cells(start_row=r_excel, end_row=r_excel, start_column=c_start, end_column=c_end)
                        except: pass # Ya fusionado o error
                    
                    cell = ws.cell(row=r_excel, column=c_start)
                    cell.fill = PatternFill("solid", fgColor=hex_color)
            
            elif val == "INFRA": # Si INFRA viene explícito en columnas
                 cell = ws.cell(row=r_excel, column=j+1)
                 cell.fill = PatternFill("solid", fgColor="000000")
                 cell.font = Font(color="FFFFFF")

    # 6. ESTILOS ESPECÍFICOS (Headers, Info Cols, AZ, Título)
    
    # Headers (Fila 2) - Color FBE2D5
    fill_head = PatternFill("solid", fgColor="FBE2D5")
    for c in range(1, max_col + 1):
        ws.cell(row=2, column=c).fill = fill_head

    # Columnas de Info (# S. TOTAL, # S. ZONA) - Color FBE2D5
    # Son las columnas 1 y 2 en el layout final
    for r in range(3, max_row + 1):
        ws.cell(row=r, column=1).fill = fill_head
        ws.cell(row=r, column=2).fill = fill_head
        # Columna SERVIDOR (3) en Negrita
        ws.cell(row=r, column=3).font = Font(name="Aptos Narrow", size=11, bold=True)

    # Columna AZ (Vertical y Color B5E6A2)
    # Necesitamos encontrar el índice de la columna AZ
    try:
        col_az_idx = df.columns.get_loc('AZ') + 1 # 1-based
        fill_az = PatternFill("solid", fgColor="B5E6A2")
        font_az = Font(name="Aptos Narrow", size=14, bold=True)
        align_az = Alignment(horizontal="center", vertical="center", textRotation=90)
        
        # Lógica de fusión vertical para AZ
        df_az_str = df['AZ'].astype(str)
        
        # Iteramos para encontrar bloques continuos de AZ
        start_row = 0
        current_val = None
        
        # Recorremos df para identificar bloques (incluyendo filas vacias como separadores)
        blocks = []
        for i, val in enumerate(df_az_str):
            if pd.isna(df.iloc[i, 0]) and pd.isna(df.iloc[i, 1]): # Fila separadora (INFRA)
                if current_val is not None:
                    blocks.append((start_row, i - 1, current_val))
                    current_val = None
                continue
                
            if val != current_val:
                if current_val is not None:
                    blocks.append((start_row, i - 1, current_val))
                current_val = val
                start_row = i
        
        if current_val is not None:
            blocks.append((start_row, len(df) - 1, current_val))

        for start_idx, end_idx, val in blocks:
            # Mapear a filas excel (start_idx + 3)
            r_start = start_idx + 3
            r_end = end_idx + 3
            
            # Aplicar estilo a todo el bloque antes de fusionar
            for r in range(r_start, r_end + 1):
                c = ws.cell(row=r, column=col_az_idx)
                c.fill = fill_az
                c.font = font_az
                c.alignment = align_az
            
            # Fusionar
            if r_end > r_start:
                ws.merge_cells(start_row=r_start, end_row=r_end, start_column=col_az_idx, end_column=col_az_idx)

    except Exception as e:
        print(f"Error estilizando AZ: {e}")

    # 7. Título (Fila 1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    tit = ws.cell(1, 1, value=f"TETRIS OPTIMIZADO - {sitio}")
    tit.font = Font(name="Aptos Narrow", size=16, bold=True, color="FF0000")
    tit.fill = PatternFill("solid", fgColor="CAEDFB")
    tit.alignment = Alignment(horizontal="center", vertical="center")

    # Guardar en bytes
    with io.BytesIO() as bio:
        wb.save(bio)
        return bio.getvalue()